#!/usr/bin/env python3
"""Convert the Chemical Inventory & EH&S Tracker workbook's Inventory sheet
into data/chemicals.json for the lookup app.

Usage:
    python3 scripts/convert_inventory.py [path/to/workbook.xlsx]

Re-run this whenever the source workbook (data/Chemical_Inventory_EHS_Tracker_1.0.xlsx,
or the sheet EH&S drops in its place) is updated, then commit the regenerated
data/chemicals.json.
"""
import json
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_WORKBOOK = REPO_ROOT / "data" / "Chemical_Inventory_EHS_Tracker_1.0.xlsx"
OUTPUT_PATH = REPO_ROOT / "data" / "chemicals.json"

# Maps workbook column headers (row 3 of the Inventory sheet) to the JSON
# field names used by the app.
FIELD_MAP = {
    "Inventory ID": "id",
    "Chemical Name": "name",
    "CAS Number": "cas",
    "Manufacturer / Supplier": "manufacturer",
    "Catalog / Product #": "catalogNumber",
    "Physical State": "physicalState",
    "Container Size": "containerSize",
    "Unit": "unit",
    "# Containers": "containerCount",
    "Total Quantity": "totalQuantity",
    "Building": "building",
    "Room": "room",
    "Storage Location": "storageLocation",
    "Storage Segregation Group": "storageSegregationGroup",
    "Primary Hazard Class": "primaryHazardClass",
    "GHS Signal Word": "ghsSignalWord",
    "Particularly Hazardous (Y/N)": "particularlyHazardous",
    "Special Regulatory Flag": "specialRegulatoryFlag",
    "SDS on File": "sdsOnFile",
    "SDS Link / Location": "sdsLink",
    "Responsible Person / PI": "responsiblePerson",
    "Date Received": "dateReceived",
    "Date Opened": "dateOpened",
    "Expiration Date": "expirationDate",
    "Days to Expiration": "daysToExpiration",
    "Expiration Status": "expirationStatus",
    "Status": "status",
    "Disposal Date": "disposalDate",
    "Notes": "notes",
    "Storage Flag": "storageFlag",
    "GHS Codes (SDS)": "ghsCodes",
}

HEADER_ROW = 3
DATA_START_ROW = 4


def clean_value(value):
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat().split("T")[0]
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def convert(workbook_path: Path) -> list:
    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    ws = wb["Inventory"]

    headers = [cell.value for cell in ws[HEADER_ROW]]
    columns = []
    for idx, header in enumerate(headers):
        field = FIELD_MAP.get(header)
        if field:
            columns.append((idx, field))
        elif header is not None:
            print(f"warning: unmapped column '{header}' at index {idx}", file=sys.stderr)

    records = []
    for row in ws.iter_rows(min_row=DATA_START_ROW, values_only=True):
        if not any(cell is not None for cell in row):
            continue
        record = {field: clean_value(row[idx]) for idx, field in columns}
        if not record.get("name") and not record.get("id"):
            continue
        records.append(record)

    return records


def main():
    workbook_path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_WORKBOOK
    if not workbook_path.exists():
        print(f"error: workbook not found at {workbook_path}", file=sys.stderr)
        sys.exit(1)

    records = convert(workbook_path)
    OUTPUT_PATH.write_text(json.dumps(records, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"wrote {len(records)} chemicals to {OUTPUT_PATH.relative_to(REPO_ROOT)}")


if __name__ == "__main__":
    main()
